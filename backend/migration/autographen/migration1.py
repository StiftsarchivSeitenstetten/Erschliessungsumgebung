"""Read-only dry run for the frozen Autographen 9.6 migration-1 workbook."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook

from backend.modules import get_module
from backend.records.generic_write import render_record_content
from backend.records.runtime import RecordRuntime, RecordValidationError


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_DIR = ROOT / "migration-work" / "autographen_9_6_migration1"
EXPECTED_SOURCE_SHA256 = "ce4aeecb20bb55b80aa67694a0333f6ad9d6d94597f46572546a1bf776914279"
SHEET_NAME = "Autographen"
MODULE_KEY = "autographen_9_6"
PROVENANCE_PERSON = "Maria Prüller"

REQUIRED_COLUMNS = (
    "Signatur",
    "Kategorie",
    "Schreiber",
    "Empfaenger",
    "Datierung",
    "Ort",
    "Regest",
    "Bemerkungen",
    "Olim",
    "InterneBemerkung",
    "ErfasstVon",
    "ErfasstAm",
    "GeaendertAm",
)
IGNORED_EXPORT_COLUMNS = ("Exportiert", "ExportiertAm", "ExportHinweis")
CATEGORY_MAPPING = {
    "Brief": ("brief", "original"),
    "Karte": ("karte", "original"),
    "Visitenkarte": ("visitenkarte", "original"),
    "Partezettel": ("partezettel", "original"),
    "Foto": ("foto", "original"),
    "Sterbebildchen": ("sterbebildchen", "original"),
    "Ausweis": ("ausweis", "original"),
    "Fragment": ("unbekannt", "fragment"),
}
SPECIAL_CASE_RULES = {
    "9.6.19": ["Leere Kategorie wird unbekannt/original."],
    "9.6.92": ["Unsicherheit (?) wird als Datumshinweis übernommen."],
    "9.6.102b": ["Empfänger Polizei-Ministerium wird als organization übernommen."],
    "9.6.110": ["Datierung und Ort werden vor der Transformation getauscht."],
    "9.6.111": ["Leere Kategorie wird unbekannt/original."],
    "9.6.144": ["Empfänger Stift Seitenstetten wird als organization übernommen."],
    "9.6.200": ["Unsicherheit (?) wird als Datumshinweis übernommen."],
    "9.6.201": ["Unsicherheit (?) wird als Datumshinweis übernommen."],
    "9.6.222": ["Zusätzliche Null in 188803003 wird entfernt; Hinweis ? bleibt erhalten."],
    "9.6.246": ["Unsicherheit (vermutlich) wird als Datumshinweis übernommen."],
    "9.6.247": ["Unsicherheit (?) wird als Datumshinweis übernommen."],
}
TARGET_SOURCE_COLUMNS = {
    "erschliessung.dokumenttyp": "Kategorie",
    "erschliessung.erhaltungsform": "Kategorie",
    "erschliessung.beteiligte": "Schreiber/Empfaenger",
    "datierung": "Datierung",
    "erschliessung.ort": "Ort",
    "erschliessung.regest": "Regest",
    "erschliessung.bemerkungen": "Bemerkungen",
    "erschliessung.altsignatur": "Olim",
    "erschliessung.interne_bemerkung": "InterneBemerkung",
    "technik.erstellt_am": "ErfasstAm",
    "technik.geaendert_am": "GeaendertAm",
}


class MigrationError(ValueError):
    """A source or transformation error that makes the dry run unsuccessful."""


@dataclass(frozen=True)
class SourceRow:
    source_row: int
    values: dict[str, Any]

    @property
    def signature(self) -> str:
        value = normalize_legacy_value(self.values.get("Signatur"))
        return value or ""


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_source_hash(path: Path, expected_sha256: str = EXPECTED_SOURCE_SHA256) -> str:
    actual = file_sha256(path)
    if actual != expected_sha256:
        raise MigrationError(f"SHA-256 stimmt nicht: erwartet {expected_sha256}, erhalten {actual}.")
    return actual


def normalize_legacy_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        text = str(value)
    elif isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text or text == "14":
        return None
    return text


def expected_signatures() -> tuple[str, ...]:
    signatures: list[str] = []
    for number in range(1, 263):
        signatures.append(f"9.6.{number}")
        if number == 102:
            signatures.extend(("9.6.102a", "9.6.102b"))
    return tuple(signatures)


def parse_signature(value: Any, bestand: str = "9.6") -> tuple[int, str | None, str]:
    text = normalize_legacy_value(value)
    match = re.fullmatch(rf"{re.escape(bestand)}\.([1-9][0-9]*)([a-z]?)", text or "")
    if not match:
        raise MigrationError(f"Ungültige Signatur: {text!r}.")
    number = int(match.group(1))
    suffix = match.group(2) or None
    return number, suffix, text


def signature_sort_key(value: Any, bestand: str = "9.6") -> tuple[int, str]:
    number, suffix, _ = parse_signature(value, bestand)
    return number, suffix or ""


def validate_signature_selection(
    rows: Iterable[SourceRow],
    required: Iterable[str] | None = None,
    bestand: str = "9.6",
) -> list[SourceRow]:
    selected = list(rows)
    signatures = [row.signature for row in selected]
    duplicates = sorted(signature for signature, count in Counter(signatures).items() if count > 1)
    expected = set(required or expected_signatures())
    actual = set(signatures)
    missing = sorted(expected - actual, key=lambda item: signature_sort_key(item, bestand))
    unexpected = sorted(actual - expected, key=lambda item: signature_sort_key(item, bestand))
    if duplicates or missing or unexpected or len(selected) != len(expected):
        details = []
        if duplicates:
            details.append(f"doppelt: {', '.join(duplicates)}")
        if missing:
            details.append(f"fehlend: {', '.join(missing)}")
        if unexpected:
            details.append(f"unerwartet: {', '.join(unexpected)}")
        if len(selected) != len(expected):
            details.append(f"Zeilen: {len(selected)}, erwartet: {len(expected)}")
        raise MigrationError("Signaturmenge für Migration 1 ist ungültig (" + "; ".join(details) + ").")
    return sorted(selected, key=lambda row: signature_sort_key(row.signature, bestand))


def _find_header_row(worksheet) -> tuple[int, list[str]]:
    required = set(REQUIRED_COLUMNS)
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1):
        headers = [str(value).strip() if value is not None else "" for value in cells]
        if required.issubset(headers):
            return row_number, headers
    raise MigrationError(f"Arbeitsblatt {SHEET_NAME!r} enthält nicht alle erforderlichen Spalten.")


def read_source_rows(source: Path) -> list[SourceRow]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise MigrationError(f"Arbeitsblatt {SHEET_NAME!r} fehlt.")
        worksheet = workbook[SHEET_NAME]
        header_row, headers = _find_header_row(worksheet)
        missing = sorted(set(REQUIRED_COLUMNS) - set(headers))
        if missing:
            raise MigrationError("Erforderliche Spalten fehlen: " + ", ".join(missing))
        rows: list[SourceRow] = []
        for source_row, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            row_values = {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header}
            if not any(value is not None and str(value).strip() for value in row_values.values()):
                continue
            rows.append(SourceRow(source_row=source_row, values=row_values))
        return rows
    finally:
        workbook.close()


def category_terms(value: Any) -> tuple[str, str]:
    category = normalize_legacy_value(value)
    if category is None:
        return "unbekannt", "original"
    try:
        return CATEGORY_MAPPING[category]
    except KeyError as exc:
        raise MigrationError(f"Unbekannte Kategorie: {category!r}.") from exc


def _date_parts(text: str) -> tuple[dict[str, int], str]:
    if re.fullmatch(r"[0-9]{4}", text):
        return {"year": int(text)}, "year_only"
    if not re.fullmatch(r"[0-9]{8}", text):
        raise MigrationError(f"Datierung ist nicht interpretierbar: {text!r}.")
    year, month, day = int(text[:4]), int(text[4:6]), int(text[6:8])
    if month == 99 and day == 99:
        return {"year": year}, "year_only"
    if day == 99 and 1 <= month <= 12:
        return {"year": year, "month": month}, "year_month"
    try:
        date(year, month, day)
    except ValueError as exc:
        raise MigrationError(f"Datierung ist ungültig: {text!r}.") from exc
    return {"year": year, "month": month, "day": day}, "full_date"


def transform_date(value: Any, signature: str) -> tuple[dict[str, Any] | None, str]:
    text = normalize_legacy_value(value)
    if text is None:
        return None, "undated"
    hint = None
    if signature == "9.6.222" and text == "188803003 (?)":
        text = "18880303 (?)"
    uncertain = re.fullmatch(r"([0-9]{4}|[0-9]{8}) \((\?|vermutlich)\)", text)
    if uncertain:
        text, hint = uncertain.group(1), uncertain.group(2)
    parts, granularity = _date_parts(text)
    transformed: dict[str, Any] = {"from": parts}
    if hint:
        transformed["hinweis"] = hint
    return transformed, granularity


def normalize_timestamp(value: Any) -> str | None:
    if value is None or normalize_legacy_value(value) is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    text = normalize_legacy_value(value)
    if text is None:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        try:
            return date.fromisoformat(text).isoformat()
        except ValueError as exc:
            raise MigrationError(f"Zeitstempel ist nicht interpretierbar: {text!r}.") from exc
    return parsed.isoformat(timespec="seconds")


def _term_ref(term_id: str, vocabulary_id: str) -> dict[str, str]:
    return {"id": term_id, "vocabulary_id": vocabulary_id}


def _participant(name: str, agent_type: str, role_id: str, role_vocabulary: str) -> dict[str, Any]:
    return {
        "agent": {"type": agent_type, "name": name},
        "rolle": _term_ref(role_id, role_vocabulary),
    }


def record_id_for_position(position: int, module) -> str:
    prefix = str(module.id_strategy["prefix"])
    width = int(module.id_strategy["width"])
    return f"{prefix}{position:0{width}d}"


def transform_row(row: SourceRow, position: int, module) -> tuple[dict[str, Any], dict[str, Any]]:
    number, suffix, signature = parse_signature(row.values.get("Signatur"), module.signature_strategy["bestand"])
    source = dict(row.values)
    if signature == "9.6.110":
        source["Datierung"], source["Ort"] = source.get("Ort"), source.get("Datierung")

    document_type, preservation_form = category_terms(source.get("Kategorie"))
    document_field = module.get_field_by_path("erschliessung.dokumenttyp")
    preservation_field = module.get_field_by_path("erschliessung.erhaltungsform")
    participants_field = module.get_field_by_path("erschliessung.beteiligte")
    role_field = next(field for field in participants_field.item_fields if field.path == "rolle")

    participants = []
    sender = normalize_legacy_value(source.get("Schreiber"))
    recipient = normalize_legacy_value(source.get("Empfaenger"))
    if sender:
        participants.append(_participant(sender, "person", "absender", role_field.vocabulary or ""))
    if recipient:
        recipient_type = "organization" if signature in {"9.6.102b", "9.6.144"} else "person"
        participants.append(_participant(recipient, recipient_type, "empfaenger", role_field.vocabulary or ""))

    transformed_date, date_granularity = transform_date(source.get("Datierung"), signature)
    created_at = normalize_timestamp(source.get("ErfasstAm"))
    changed_at = normalize_timestamp(source.get("GeaendertAm"))
    signature_statuses = module.signature_strategy.get("status_values") or ["vergeben"]
    signature_format = (module.signature_strategy.get("partitions") or ["A"])[0]
    record = {
        "schema_version": module.create_strategy.get("server_values", {}).get("schema_version", 1),
        "id": record_id_for_position(position, module),
        "datensatz_typ": module.record_type,
        "modul": module.create_strategy.get("server_values", {}).get("modul", module.access_key),
        "signatur": {
            "bestand": module.signature_strategy["bestand"],
            "format": signature_format,
            "nummer": number,
            "zusatz": suffix,
            "anzeige": signature,
            "status": signature_statuses[0],
        },
        "erschliessung": {
            "dokumenttyp": _term_ref(document_type, document_field.vocabulary or ""),
            "erhaltungsform": _term_ref(preservation_form, preservation_field.vocabulary or ""),
            "beteiligte": participants,
        },
        "technik": {
            "erstellt_am": created_at,
            "erstellt_von": PROVENANCE_PERSON,
            "geaendert_am": changed_at,
            "geaendert_von": PROVENANCE_PERSON if changed_at else None,
        },
    }
    if transformed_date:
        record["datierung"] = transformed_date
    optional_text = {
        "ort": normalize_legacy_value(source.get("Ort")),
        "regest": normalize_legacy_value(source.get("Regest")),
        "bemerkungen": normalize_legacy_value(source.get("Bemerkungen")),
        "altsignatur": normalize_legacy_value(source.get("Olim")),
        "interne_bemerkung": normalize_legacy_value(source.get("InterneBemerkung")),
    }
    if optional_text["ort"]:
        record["erschliessung"]["ort"] = {"name": optional_text["ort"]}
    for field in ("regest", "bemerkungen", "altsignatur", "interne_bemerkung"):
        if optional_text[field]:
            record["erschliessung"][field] = optional_text[field]
    facts = {
        "signature": signature,
        "source_row": row.source_row,
        "sender": sender,
        "recipient": recipient,
        "recipient_type": "organization" if recipient and signature in {"9.6.102b", "9.6.144"} else ("person" if recipient else None),
        "date_granularity": date_granularity,
        "date_hint": transformed_date.get("hinweis") if transformed_date else None,
        **optional_text,
    }
    return record, facts


def _value_at_path(record: dict[str, Any], path: str) -> Any:
    value: Any = record
    for segment in path.split("."):
        if not isinstance(value, dict) or segment not in value:
            return None
        value = value[segment]
    return value


def _source_value_for_path(source_values: dict[str, Any] | None, path: str) -> Any:
    if not source_values:
        return None
    source_column = TARGET_SOURCE_COLUMNS.get(path)
    if source_column == "Schreiber/Empfaenger":
        return {
            "Schreiber": _json_value(source_values.get("Schreiber")),
            "Empfaenger": _json_value(source_values.get("Empfaenger")),
        }
    return _json_value(source_values.get(source_column)) if source_column else None


def validate_record(
    record: dict[str, Any],
    runtime: RecordRuntime,
    signature: str,
    source_values: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    validators = (
        ("schema", runtime.validate),
        ("vocabulary", runtime.validate_vocabulary_references),
    )
    for error_type, validator in validators:
        try:
            validator(record)
        except RecordValidationError as exc:
            for message in exc.errors:
                field, separator, detail = message.partition(": ")
                field = field if separator else "<record>"
                errors.append({
                    "signatur": signature,
                    "error_type": error_type,
                    "field": field,
                    "source_value": _source_value_for_path(source_values, field),
                    "generated_value": _json_value(_value_at_path(record, field)),
                    "message": message,
                })
    return errors


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return value


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(_json_value(value), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _matrix_row(record: dict[str, Any], facts: dict[str, Any], status: str) -> dict[str, Any]:
    date_value = record.get("datierung")
    return {
        "Signatur": facts["signature"],
        "technische ID": record["id"],
        "Dokumenttyp": record["erschliessung"]["dokumenttyp"]["id"],
        "Erhaltungsform": record["erschliessung"]["erhaltungsform"]["id"],
        "Schreiber": facts["sender"] or "",
        "Empfänger": facts["recipient"] or "",
        "Agententyp Empfänger": facts["recipient_type"] or "",
        "strukturierte Datierung": json.dumps(date_value, ensure_ascii=False, separators=(",", ":")) if date_value else "",
        "Datumshinweis": facts["date_hint"] or "",
        "Ort": facts["ort"] or "",
        "Validierungsstatus": status,
    }


def _build_summary(
    source: Path,
    source_hash: str,
    source_rows: int,
    selected_records: int,
    records: list[dict[str, Any]],
    facts_by_signature: dict[str, dict[str, Any]],
    errors: list[dict[str, Any]],
    warnings: list[dict[str, Any]],
) -> dict[str, Any]:
    documents = Counter(record["erschliessung"]["dokumenttyp"]["id"] for record in records)
    preservation = Counter(record["erschliessung"]["erhaltungsform"]["id"] for record in records)
    roles = Counter()
    agent_types = Counter()
    for record in records:
        for participant in record["erschliessung"]["beteiligte"]:
            roles[participant["rolle"]["id"]] += 1
            agent_types[participant["agent"]["type"]] += 1
    date_stats = Counter(
        facts_by_signature[record["signatur"]["anzeige"]]["date_granularity"]
        for record in records
    )
    date_stats["with_hint"] = sum(bool(facts["date_hint"]) for facts in facts_by_signature.values())
    completeness = {
        field: sum(field in record["erschliessung"] for record in records)
        for field in ("ort", "regest", "bemerkungen", "altsignatur", "interne_bemerkung")
    }
    technical = {
        "erstellt_am": sum(bool(record["technik"]["erstellt_am"]) for record in records),
        "geaendert_am": sum(bool(record["technik"]["geaendert_am"]) for record in records),
    }
    return {
        "source_filename": source.name,
        "source_sha256": source_hash,
        "source_rows": source_rows,
        "selected_records": selected_records,
        "generated_records": len(records),
        "validation_errors": len(errors),
        "schema_errors": sum(error.get("error_type") == "schema" for error in errors),
        "vocabulary_errors": sum(error.get("error_type") == "vocabulary" for error in errors),
        "warnings": len(warnings),
        "warning_details": warnings,
        "first_signature": records[0]["signatur"]["anzeige"] if records else None,
        "last_signature": records[-1]["signatur"]["anzeige"] if records else None,
        "first_record_id": records[0]["id"] if records else None,
        "last_record_id": records[-1]["id"] if records else None,
        "document_type_distribution": dict(sorted(documents.items())),
        "preservation_form_distribution": dict(sorted(preservation.items())),
        "participants": {
            "total": sum(roles.values()),
            "by_role": dict(sorted(roles.items())),
            "by_agent_type": dict(sorted(agent_types.items())),
        },
        "date_statistics": dict(date_stats),
        "field_completeness": completeness,
        "technical_metadata_statistics": technical,
        "special_corrections": [
            {"signatur": signature, "rules": rules}
            for signature, rules in SPECIAL_CASE_RULES.items()
        ],
    }


def run_dry_run(
    source: Path,
    output_dir: Path,
    *,
    expected_sha256: str = EXPECTED_SOURCE_SHA256,
    required_signatures: Iterable[str] | None = None,
) -> dict[str, Any]:
    source = source.resolve()
    output_dir = output_dir.resolve()
    source_hash = verify_source_hash(source, expected_sha256)
    module = get_module(MODULE_KEY)
    rows = read_source_rows(source)
    sorted_rows = validate_signature_selection(
        rows,
        required=required_signatures,
        bestand=module.signature_strategy["bestand"],
    )
    if output_dir.exists() and any(output_dir.iterdir()):
        raise MigrationError(f"Ausgabeordner ist nicht leer: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)
    previews = output_dir / "preview"
    previews.mkdir()

    runtime = RecordRuntime(module)
    records: list[dict[str, Any]] = []
    facts_by_signature: dict[str, dict[str, Any]] = {}
    matrix: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    special_cases: list[dict[str, Any]] = []

    for position, row in enumerate(sorted_rows, start=1):
        signature = row.signature
        try:
            record, facts = transform_row(row, position, module)
        except MigrationError as exc:
            errors.append({
                "signatur": signature or None,
                "field": "<transformation>",
                "source_value": _json_value(row.values),
                "generated_value": None,
                "message": str(exc),
            })
            continue
        records.append(record)
        facts_by_signature[signature] = facts
        record_errors = validate_record(record, runtime, signature, row.values)
        errors.extend(record_errors)
        status = "OK" if not record_errors else f"FEHLER ({len(record_errors)})"
        matrix.append(_matrix_row(record, facts, status))
        (previews / f"{record['id']}.md").write_text(render_record_content(record), encoding="utf-8")
        if signature in SPECIAL_CASE_RULES:
            special_cases.append({
                "signatur": signature,
                "rules": SPECIAL_CASE_RULES[signature],
                "source": _json_value({field: row.values.get(field) for field in REQUIRED_COLUMNS}),
                "transformed": record,
                "validation_status": status,
            })

    duplicate_ids = sorted(record_id for record_id, count in Counter(record["id"] for record in records).items() if count > 1)
    if duplicate_ids:
        errors.append({
            "signatur": None,
            "field": "id",
            "source_value": None,
            "generated_value": duplicate_ids,
            "message": "Doppelte technische IDs.",
        })
    summary = _build_summary(
        source, source_hash, len(rows), len(sorted_rows), records, facts_by_signature, errors, warnings,
    )
    _write_json(output_dir / "summary.json", summary)
    _write_json(output_dir / "special_cases.json", special_cases)
    _write_json(output_dir / "errors.json", errors)
    with (output_dir / "migration_matrix.csv").open("w", encoding="utf-8", newline="") as handle:
        fieldnames = list(matrix[0]) if matrix else [
            "Signatur", "technische ID", "Dokumenttyp", "Erhaltungsform", "Schreiber",
            "Empfänger", "Agententyp Empfänger", "strukturierte Datierung", "Datumshinweis",
            "Ort", "Validierungsstatus",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(matrix)
    return {"summary": summary, "errors": errors, "warnings": warnings, "output_dir": str(output_dir)}


def print_report(result: dict[str, Any]) -> None:
    summary = result["summary"]
    print(f"Quelldatei: {summary['source_filename']}")
    print(f"SHA-256: OK ({summary['source_sha256']})")
    print(f"Gelesene Excel-Zeilen: {summary['source_rows']}")
    print(f"Ausgewählte Datensätze: {summary['selected_records']}")
    print(f"Zielrecords: {summary['generated_records']}")
    print(f"Signaturen: {summary['first_signature']} bis {summary['last_signature']}")
    print(f"Technische IDs: {summary['first_record_id']} bis {summary['last_record_id']}")
    print(f"Schemafehler: {summary['schema_errors']}")
    print(f"Vokabularfehler: {summary['vocabulary_errors']}")
    print(f"Validierungsfehler gesamt: {summary['validation_errors']}")
    print(f"Warnungen: {summary['warnings']}")
    print(f"Sonderkorrekturen: {len(summary['special_corrections'])}")
    print(f"Ausgabeordner: {result['output_dir']}")
    if summary["validation_errors"] == 0 and summary["generated_records"] == summary["selected_records"]:
        print("DRY RUN ERFOLGREICH – ES WURDEN KEINE PRODUKTIVDATEN VERÄNDERT.")
    else:
        print("DRY RUN FEHLGESCHLAGEN – ES WURDEN KEINE PRODUKTIVDATEN VERÄNDERT.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Read-only Dry Run für Autographen 9.6, Migration 1")
    parser.add_argument("--source", required=True, type=Path, help="Eingefrorene XLSM-Quelldatei")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Leerer lokaler Prüfartefakt-Ordner")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        result = run_dry_run(args.source, args.output_dir)
    except (MigrationError, OSError) as exc:
        print(f"DRY RUN ABGEBROCHEN: {exc}")
        print("ES WURDEN KEINE PRODUKTIVDATEN VERÄNDERT.")
        return 2
    print_report(result)
    return 0 if not result["errors"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
