"""Analyse and dry-run migration for legacy photo Excel workbooks."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, time
from hashlib import sha256
import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel

from scripts.foto_core import build_signature, render_photo_markdown, validate_record, validate_record_schema


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_DIR = ROOT / "migration-work" / "foto-dry-run"
TABLE_NAME = "tbl_Fotos"
FORMATS = ("A", "B", "C", "D", "E", "F")
FIELD_NAMES = [
    "Format",
    "Nummer",
    "Signatur",
    "Titel",
    "Beschriftung",
    "Beschreibung",
    "Datierung",
    "DargestelltePersonen",
    "Herkunft",
    "Sammler",
    "Fotograf",
    "Rechteinhaber",
    "Orte",
    "Schlagworte",
    "Altsignatur",
    "InterneBemerkung",
    "ZuKlaeren",
    "ExportBereit",
    "Exportiert",
    "ExportiertAm",
    "ExportHinweis",
]


@dataclass(frozen=True)
class SourceRow:
    source_row: int
    order: int
    values: dict[str, Any]
    datierung_cell: Cell


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def scalar(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def parse_number(value: Any) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float) and value.is_integer():
        return int(value) if value > 0 else None
    text = scalar(value)
    if text and re.fullmatch(r"[0-9]+", text):
        number = int(text)
        return number if number > 0 else None
    return None


def normalize_format(value: Any) -> str | None:
    text = scalar(value)
    if not text:
        return None
    return text.upper()


def split_semicolon(value: Any) -> list[str]:
    text = scalar(value)
    if not text:
        return []
    return [part.strip() for part in text.split(";") if part.strip()]


def row_fingerprint(row: SourceRow) -> str:
    payload = json.dumps({key: scalar(value) for key, value in row.values.items()}, ensure_ascii=False, sort_keys=True)
    return sha256(payload.encode("utf-8")).hexdigest()


def read_table(source: Path, table_name: str = TABLE_NAME) -> list[SourceRow]:
    wb = load_workbook(source, read_only=False, data_only=True, keep_vba=False)
    ws = wb["Fotos"]
    table = ws.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
    if headers != FIELD_NAMES:
        raise ValueError(f"Unerwartete Spalten in {table_name}: {headers}")
    datierung_index = headers.index("Datierung") + min_col
    rows: list[SourceRow] = []
    for order, row_idx in enumerate(range(min_row + 1, max_row + 1), start=1):
        values = {header: ws.cell(row_idx, min_col + offset).value for offset, header in enumerate(headers)}
        rows.append(SourceRow(source_row=row_idx, order=order, values=values, datierung_cell=ws.cell(row_idx, datierung_index)))
    return rows


def parse_signature(text: str | None) -> dict[str, Any] | None:
    if not text:
        return None
    match = re.fullmatch(r"9\.4\.2\.([A-F])\.([1-9][0-9]*)([a-z]?)", text.strip())
    if not match:
        return None
    return {"format": match.group(1), "nummer": int(match.group(2)), "zusatz": match.group(3) or None}


def display_signature(format_code: str, number: int, suffix: str | None = None) -> str:
    return f"9.4.2.{format_code}.{number}{suffix or ''}"


def suffix_for_index(index: int) -> str:
    alphabet = "abcdefghijklmnopqrstuvwxyz"
    if index < len(alphabet):
        return alphabet[index]
    return alphabet[index // len(alphabet) - 1] + alphabet[index % len(alphabet)]


def classify_date(cell: Cell) -> dict[str, Any]:
    value = cell.value
    if value in (None, ""):
        return {"category": "leer", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": None, "problem": None}
    if isinstance(value, datetime):
        return {
            "category": "Excel-Datum",
            "datierung": {"jahr": value.year, "monat": value.month, "tag": value.day},
            "original": value.isoformat(),
            "problem": None,
        }
    if isinstance(value, (int, float)) and cell.is_date:
        dt = from_excel(value)
        return {
            "category": "Excel-Datum",
            "datierung": {"jahr": dt.year, "monat": dt.month, "tag": dt.day},
            "original": str(value),
            "problem": None,
        }
    text = scalar(value)
    if not text:
        return {"category": "leer", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": None, "problem": None}
    normalized = text.strip()
    compact = re.fullmatch(r"([0-9]{4})([0-9]{2})([0-9]{2})", normalized)
    dotted = re.fullmatch(r"([0-9]{1,2})\.([0-9]{1,2})\.([0-9]{4})", normalized)
    if compact or dotted:
        if compact:
            year, month, day = int(compact.group(1)), int(compact.group(2)), int(compact.group(3))
        else:
            day, month, year = int(dotted.group(1)), int(dotted.group(2)), int(dotted.group(3))
        if year < 1:
            return {"category": "offensichtlich problematisch", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "Jahr 0 ist nicht plausibel"}
        if month in (0, 99) and day in (0, 99):
            return {"category": "sicher nur Jahr", "datierung": {"jahr": year, "monat": None, "tag": None}, "original": normalized, "problem": None}
        if day in (0, 99) and 1 <= month <= 12:
            return {"category": "sicher Jahr/Monat", "datierung": {"jahr": year, "monat": month, "tag": None}, "original": normalized, "problem": None}
        try:
            datetime(year, month, day)
        except ValueError:
            return {"category": "offensichtlich problematisch", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "ungueltige Tages-/Monatskombination"}
        return {"category": "sicher vollständig konvertiert", "datierung": {"jahr": year, "monat": month, "tag": day}, "original": normalized, "problem": None}
    if re.fullmatch(r"[0-9]{4}", normalized):
        year = int(normalized)
        return {"category": "sicher nur Jahr", "datierung": {"jahr": year, "monat": None, "tag": None}, "original": normalized, "problem": None}
    if re.fullmatch(r"[0-9]{8}-[0-9]{8}", normalized) or " - " in normalized or " bis " in normalized.lower():
        return {"category": "Zeitraum", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "Zeitraum nicht automatisch strukturiert"}
    if re.search(r"\b(um|ca\.?|circa|sommer|winter|fruehjahr|frühjahr|herbst)\b", normalized.lower()):
        return {"category": "unscharfe Textdatierung", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "unscharfe Textdatierung"}
    if re.fullmatch(r"[0-9]{5}", normalized):
        return {"category": "nicht interpretierbar", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "fuenfstellige Zahl ohne Zelltyp nicht eindeutig"}
    return {"category": "nicht interpretierbar", "datierung": {"jahr": None, "monat": None, "tag": None}, "original": normalized, "problem": "keine sichere Datierungsregel"}


def analyze_separators(rows: list[SourceRow], field: str) -> dict[str, Any]:
    counts = Counter()
    non_empty = 0
    samples = []
    for row in rows:
        text = scalar(row.values.get(field))
        if not text:
            continue
        non_empty += 1
        for char in (";", ",", "/", "|", "\n"):
            if char in text:
                counts[char] += 1
        if len(samples) < 20:
            samples.append(text)
    return {"non_empty": non_empty, "separator_counts": dict(counts), "samples": samples}


def signature_analysis(rows: list[SourceRow]) -> dict[str, Any]:
    format_counts = {format_code: 0 for format_code in FORMATS}
    numbers_by_format: dict[str, list[int]] = {format_code: [] for format_code in FORMATS}
    missing_numbers = []
    missing_signatures = []
    invalid_signatures = []
    mismatches = []
    canonical_keys: list[tuple[str, int, str, int]] = []
    for row in rows:
        format_code = normalize_format(row.values.get("Format"))
        number = parse_number(row.values.get("Nummer"))
        signature = scalar(row.values.get("Signatur"))
        parsed_signature = parse_signature(signature)
        if format_code in format_counts:
            format_counts[format_code] += 1
        if format_code in numbers_by_format and number is not None:
            numbers_by_format[format_code].append(number)
            canonical_keys.append((format_code, number, display_signature(format_code, number), row.order))
        if number is None:
            missing_numbers.append(row.source_row)
        if not signature:
            missing_signatures.append(row.source_row)
        elif parsed_signature is None:
            invalid_signatures.append({"source_row": row.source_row, "signatur": signature})
        elif format_code != parsed_signature["format"] or number != parsed_signature["nummer"]:
            mismatches.append({
                "source_row": row.source_row,
                "format": format_code,
                "nummer": number,
                "signatur": signature,
                "parsed": parsed_signature,
            })
    groups = defaultdict(list)
    for format_code, number, signature, order in canonical_keys:
        groups[(format_code, number, signature)].append(order)
    duplicates = []
    for (format_code, number, signature), orders in sorted(groups.items(), key=lambda item: item[0]):
        if len(orders) > 1:
            duplicates.append({"format": format_code, "nummer": number, "original_signatur": signature, "orders": orders})
    return {
        "format_counts": format_counts,
        "numbers_by_format": numbers_by_format,
        "missing_numbers": missing_numbers,
        "missing_signatures": missing_signatures,
        "invalid_signatures": invalid_signatures,
        "mismatches": mismatches,
        "duplicates": duplicates,
    }


def build_manifest_and_records(rows: list[SourceRow], sig: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    duplicate_lookup: dict[tuple[str, int], list[int]] = {}
    for group in sig["duplicates"]:
        duplicate_lookup[(group["format"], group["nummer"])] = group["orders"]
    manifest = []
    records = []
    conflicts = []
    for row in rows:
        target_id = f"foto-{row.order:06d}"
        format_code = normalize_format(row.values.get("Format"))
        number = parse_number(row.values.get("Nummer"))
        source_signature = scalar(row.values.get("Signatur"))
        status = "ready"
        suffix = None
        if format_code not in FORMATS or number is None or not source_signature:
            status = "conflict_missing_signature_parts"
        elif (format_code, number) in duplicate_lookup:
            orders = duplicate_lookup[(format_code, number)]
            suffix = suffix_for_index(orders.index(row.order))
        date_info = classify_date(row.datierung_cell)
        alts = split_semicolon(row.values.get("Altsignatur"))
        if suffix and source_signature and source_signature not in alts:
            alts.append(source_signature)
        source = {
            "source_table": TABLE_NAME,
            "source_row": row.source_row,
            "source_order": row.order,
            "source_signatur": source_signature,
            "target_id": target_id,
            "row_fingerprint": row_fingerprint(row),
            "status": status,
        }
        manifest.append(source)
        if status != "ready":
            conflicts.append({**source, "reason": status})
            continue
        signature = build_signature(format_code, number, status="vergeben")
        if suffix:
            signature["zusatz"] = suffix
            signature["anzeige"] = display_signature(format_code, number, suffix)
        record = {
            "schema_version": 1,
            "id": target_id,
            "datensatz_typ": "foto",
            "modul": "papierabzuege_9_4_2",
            "signatur": signature,
            "erschliessung": {
                "titel": scalar(row.values.get("Titel")),
                "beschriftung": scalar(row.values.get("Beschriftung")),
                "beschreibung": scalar(row.values.get("Beschreibung")),
                "dargestellte_personen": [{"name": name, "hinweis": None} for name in split_semicolon(row.values.get("DargestelltePersonen"))],
                "herkunft": scalar(row.values.get("Herkunft")),
                "sammler": scalar(row.values.get("Sammler")),
                "fotograf": scalar(row.values.get("Fotograf")),
                "rechteinhaber": scalar(row.values.get("Rechteinhaber")),
                "orte": split_semicolon(row.values.get("Orte")) or ([scalar(row.values.get("Orte"))] if scalar(row.values.get("Orte")) else []),
                "schlagworte": split_semicolon(row.values.get("Schlagworte")) or ([scalar(row.values.get("Schlagworte"))] if scalar(row.values.get("Schlagworte")) else []),
                "altsignaturen": alts,
                "interne_bemerkung": scalar(row.values.get("InterneBemerkung")),
            },
            "korrespondenzstueck": None,
            "datierung": {
                **date_info["datierung"],
                "anmerkung": None,
                "original": date_info["original"],
                "original_typ": "excel_altbestand" if date_info["original"] else None,
            },
            "redaktion": {"stufe": "ehrenamtlich"},
            "bearbeitung": {"status": "in_bearbeitung"},
            "publikation": {"status": "intern"},
            "technik": {"quelle": "migration_excel"},
        }
        records.append(record)
        errors = validate_record_schema(record) + validate_record(record)
        if errors:
            conflicts.append({**source, "reason": "schema_or_validation", "errors": errors})
    return manifest, records, conflicts


def analyze(source: Path, output_dir: Path = DEFAULT_OUTPUT_DIR) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = read_table(source)
    digest = file_sha256(source)
    sig = signature_analysis(rows)
    manifest, records, conflicts = build_manifest_and_records(rows, sig)
    date_stats = Counter()
    date_problems = []
    for row in rows:
        info = classify_date(row.datierung_cell)
        date_stats[info["category"]] += 1
        if info["problem"]:
            date_problems.append({"source_row": row.source_row, "value": scalar(row.values.get("Datierung")), "category": info["category"], "problem": info["problem"]})
    field_coverage = {}
    for field in FIELD_NAMES:
        field_coverage[field] = sum(1 for row in rows if scalar(row.values.get(field)) is not None)
    numbers = sig["numbers_by_format"]
    gaps = {}
    minmax = {}
    next_signature = {}
    for format_code in FORMATS:
        nums = sorted(set(numbers[format_code]))
        if nums:
            minmax[format_code] = {"min": min(nums), "max": max(nums)}
            gaps[format_code] = [number for number in range(min(nums), max(nums) + 1) if number not in set(nums)]
            next_signature[format_code] = max(nums) + 1
        else:
            minmax[format_code] = {"min": None, "max": None}
            gaps[format_code] = []
            next_signature[format_code] = 1
    duplicate_assignments = []
    rows_by_order = {row.order: row for row in rows}
    duplicate_classification = {"identical_groups": 0, "different_groups": 0, "groups": []}
    for group in sig["duplicates"]:
        fingerprints = {row_fingerprint(rows_by_order[order]) for order in group["orders"] if order in rows_by_order}
        classification = "identisch" if len(fingerprints) == 1 else "inhaltlich_unterschiedlich"
        if classification == "identisch":
            duplicate_classification["identical_groups"] += 1
        else:
            duplicate_classification["different_groups"] += 1
        duplicate_classification["groups"].append({
            "original_signatur": group["original_signatur"],
            "orders": group["orders"],
            "classification": classification,
        })
        for index, order in enumerate(group["orders"]):
            duplicate_assignments.append({
                "source_order": order,
                "original_signatur": group["original_signatur"],
                "target_signatur": display_signature(group["format"], group["nummer"], suffix_for_index(index)),
                "zusatz": suffix_for_index(index),
            })
    report = {
        "source_path": str(source),
        "source_sha256": digest,
        "table": TABLE_NAME,
        "record_count": len(rows),
        "format_counts": sig["format_counts"],
        "number_minmax": minmax,
        "highest_number": {format_code: minmax[format_code]["max"] for format_code in FORMATS},
        "missing_numbers_count": len(sig["missing_numbers"]),
        "missing_numbers_rows": sig["missing_numbers"][:500],
        "missing_signatures_count": len(sig["missing_signatures"]),
        "missing_signatures_rows": sig["missing_signatures"][:500],
        "invalid_signatures_count": len(sig["invalid_signatures"]),
        "invalid_signatures": sig["invalid_signatures"][:500],
        "signature_mismatches_count": len(sig["mismatches"]),
        "signature_mismatches": sig["mismatches"][:500],
        "duplicate_group_count": len(sig["duplicates"]),
        "duplicate_record_count": sum(len(group["orders"]) for group in sig["duplicates"]),
        "duplicate_assignments": duplicate_assignments,
        "duplicate_classification": duplicate_classification,
        "gaps": {format_code: {"count": len(values), "values": values[:500]} for format_code, values in gaps.items()},
        "field_coverage": field_coverage,
        "date_stats": dict(date_stats),
        "date_problems_count": len(date_problems),
        "date_problems": date_problems[:1000],
        "person_stats": {
            "rows_with_persons": sum(1 for row in rows if scalar(row.values.get("DargestelltePersonen"))),
            "total_person_entries": sum(len(split_semicolon(row.values.get("DargestelltePersonen"))) for row in rows),
        },
        "separator_analysis": {field: analyze_separators(rows, field) for field in ("Orte", "Schlagworte", "Altsignatur")},
        "workflow_coverage": {field: field_coverage[field] for field in ("ZuKlaeren", "ExportBereit", "Exportiert", "ExportiertAm", "ExportHinweis")},
        "conflict_count": len(conflicts),
        "recommended_state": {"next_record_id": len(rows) + 1, "next_signature_number": next_signature},
    }
    (output_dir / "analysis-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "migration-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "conflicts.json").write_text(json.dumps(conflicts, ensure_ascii=False, indent=2), encoding="utf-8")
    samples = records[:10]
    sample_dir = output_dir / "sample-yaml"
    sample_dir.mkdir(exist_ok=True)
    for record in samples:
        (sample_dir / f"{record['id']}.md").write_text(render_photo_markdown(record), encoding="utf-8")
    write_markdown_report(output_dir / "analysis-report.md", report)
    return report


def write_markdown_report(path: Path, report: dict[str, Any]) -> None:
    lines = [
        "# Foto-Migration Dry-Run",
        "",
        f"Quelle: `{report['source_path']}`",
        f"SHA-256: `{report['source_sha256']}`",
        f"Tabelle: `{report['table']}`",
        f"Datensätze: {report['record_count']}",
        "",
        "## Formatverteilung",
    ]
    for format_code, count in report["format_counts"].items():
        highest = report["highest_number"][format_code]
        lines.append(f"- {format_code}: {count} Datensätze, höchste Nummer: {highest}")
    lines += [
        "",
        "## Signaturen",
        f"- Fehlende Nummern: {report['missing_numbers_count']}",
        f"- Fehlende Signaturen: {report['missing_signatures_count']}",
        f"- Formal ungültige Signaturen: {report['invalid_signatures_count']}",
        f"- Abweichungen Format/Nummer/Signatur: {report['signature_mismatches_count']}",
        f"- Dublettengruppen: {report['duplicate_group_count']}",
        f"- Von Dubletten betroffene Datensätze: {report['duplicate_record_count']}",
        "",
        "## Datierung",
    ]
    for category, count in sorted(report["date_stats"].items()):
        lines.append(f"- {category}: {count}")
    lines += [
        "",
        "## Vorgeschlagener State",
        "```json",
        json.dumps(report["recommended_state"], ensure_ascii=False, indent=2),
        "```",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["analyse"])
    parser.add_argument("--source", required=True)
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    args = parser.parse_args(argv)
    report = analyze(Path(args.source), Path(args.output_dir))
    print(json.dumps({
        "source_sha256": report["source_sha256"],
        "record_count": report["record_count"],
        "format_counts": report["format_counts"],
        "output_dir": args.output_dir,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
