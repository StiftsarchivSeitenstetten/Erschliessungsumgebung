from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from backend.migration.fotos.analyse import (
    FIELD_NAMES,
    SourceRow,
    analyze,
    build_manifest_and_records,
    classify_date,
    split_semicolon,
)
from exports.archivis.export import yes_no_empty
from scripts.foto_core import validate_record_schema


def source_row(order, *, format_code="A", number=1, signature="9.4.2.A.1", persons="", altsignatur="", datierung=None):
    wb = Workbook()
    ws = wb.active
    cell = ws["A1"]
    cell.value = datierung
    return SourceRow(
        source_row=order + 1,
        order=order,
        datierung_cell=cell,
        values={
            "Format": format_code,
            "Nummer": number,
            "Signatur": signature,
            "Titel": None,
            "Beschriftung": f"Beschriftung {order}",
            "Beschreibung": None,
            "Datierung": datierung,
            "DargestelltePersonen": persons,
            "Herkunft": None,
            "Sammler": None,
            "Fotograf": None,
            "Rechteinhaber": None,
            "Orte": None,
            "Schlagworte": None,
            "Altsignatur": altsignatur,
            "InterneBemerkung": None,
            "ZuKlaeren": None,
            "ExportBereit": None,
            "Exportiert": None,
            "ExportiertAm": None,
            "ExportHinweis": None,
        },
    )


class FotoMigrationTest(unittest.TestCase):
    def test_persons_split_by_semicolon_without_normalization(self):
        self.assertEqual(split_semicolon("Abt Albert; Kurzwernhart, Albert; ; P. Ambros;"), ["Abt Albert", "Kurzwernhart, Albert", "P. Ambros"])

    def test_safe_date_forms_and_uncertain_text(self):
        row = source_row(1, datierung="19801225")
        self.assertEqual(classify_date(row.datierung_cell)["datierung"], {"jahr": 1980, "monat": 12, "tag": 25})
        row = source_row(1, datierung="99.12.1980")
        self.assertEqual(classify_date(row.datierung_cell)["datierung"], {"jahr": 1980, "monat": 12, "tag": None})
        row = source_row(1, datierung="1966")
        self.assertEqual(classify_date(row.datierung_cell)["datierung"], {"jahr": 1966, "monat": None, "tag": None})
        row = source_row(1, datierung="um 1890")
        self.assertEqual(classify_date(row.datierung_cell)["category"], "unscharfe Textdatierung")
        self.assertIsNone(classify_date(row.datierung_cell)["datierung"]["jahr"])
        row = source_row(1, datierung="00000000")
        self.assertEqual(classify_date(row.datierung_cell)["category"], "offensichtlich problematisch")
        self.assertIsNone(classify_date(row.datierung_cell)["datierung"]["jahr"])

    def test_excel_date_cell(self):
        row = source_row(1, datierung=datetime(1980, 12, 25))
        info = classify_date(row.datierung_cell)
        self.assertEqual(info["category"], "Excel-Datum")
        self.assertEqual(info["datierung"], {"jahr": 1980, "monat": 12, "tag": 25})

    def test_duplicate_signatures_get_suffix_for_all_members(self):
        rows = [
            source_row(1, number=5528, signature="9.4.2.A.5528", altsignatur="alt"),
            source_row(2, number=5528, signature="9.4.2.A.5528"),
            source_row(3, number=5528, signature="9.4.2.A.5528"),
        ]
        sig = {
            "duplicates": [{"format": "A", "nummer": 5528, "original_signatur": "9.4.2.A.5528", "orders": [1, 2, 3]}],
        }
        _, records, _ = build_manifest_and_records(rows, sig)
        self.assertEqual([record["signatur"]["zusatz"] for record in records], ["a", "b", "c"])
        self.assertEqual([record["signatur"]["anzeige"] for record in records], ["9.4.2.A.5528a", "9.4.2.A.5528b", "9.4.2.A.5528c"])
        self.assertEqual([record["signatur"]["nummer"] for record in records], [5528, 5528, 5528])
        self.assertIn("alt", records[0]["erschliessung"]["altsignaturen"])
        self.assertIn("9.4.2.A.5528", records[0]["erschliessung"]["altsignaturen"])
        self.assertIn("9.4.2.A.5528", records[1]["erschliessung"]["altsignaturen"])

    def test_missing_signature_is_conflict_but_gets_id(self):
        rows = [source_row(1, number=None, signature=None)]
        manifest, records, conflicts = build_manifest_and_records(rows, {"duplicates": []})
        self.assertEqual(manifest[0]["target_id"], "foto-000001")
        self.assertEqual(manifest[0]["status"], "conflict_missing_signature_parts")
        self.assertEqual(records, [])
        self.assertEqual(conflicts[0]["target_id"], "foto-000001")

    def test_legacy_korrespondenz_null_schema_and_archive_export(self):
        rows = [source_row(1)]
        _, records, _ = build_manifest_and_records(rows, {"duplicates": []})
        self.assertIsNone(records[0]["korrespondenzstueck"])
        self.assertEqual(validate_record_schema(records[0]), [])
        self.assertEqual(yes_no_empty(None), "")
        self.assertEqual(yes_no_empty(False), "Nein")
        self.assertEqual(yes_no_empty(True), "Ja")

    def test_dry_run_reproducible_ids_and_no_github_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "Fotoerfassung-test.xlsm"
            output_a = Path(tmp) / "a"
            output_b = Path(tmp) / "b"
            self._write_workbook(source)
            report_a = analyze(source, output_a)
            report_b = analyze(source, output_b)
            manifest_a = json.loads((output_a / "migration-manifest.json").read_text(encoding="utf-8"))
            manifest_b = json.loads((output_b / "migration-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual([item["target_id"] for item in manifest_a], [item["target_id"] for item in manifest_b])
            self.assertEqual(report_a["recommended_state"]["next_record_id"], 4)
            self.assertEqual(report_b["recommended_state"]["next_record_id"], 4)
            self.assertFalse((Path(tmp) / "Erschliessungsdaten").exists())

    def _write_workbook(self, path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Fotos"
        ws.append(FIELD_NAMES)
        ws.append(["A", 1, "9.4.2.A.1", None, "B1", None, "1966", "Abt Albert; P. Ambros", None, None, None, None, None, None, None, None, None, None, None, None, None])
        ws.append(["A", 1, "9.4.2.A.1", None, "B2", None, "19801225", None, None, None, None, None, None, None, None, None, None, None, None, None, None])
        ws.append(["B", None, None, None, "B3", None, "um 1890", None, None, None, None, None, None, None, None, None, None, None, None, None, None])
        table = Table(displayName="tbl_Fotos", ref="A1:U4")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        wb.save(path)


if __name__ == "__main__":
    unittest.main()
